"""Mejora 9: plantilla Excel fija para cargar notas (descarga, validación y vista previa)."""

from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import load_workbook

from . import excel
from .models import Calificacion
from .tests import GradesBase


class ExcelBase(GradesBase):
    def descargar(self, user=None, actividad=None):
        self.auth(user or self.doc1)
        return self.client.get(
            reverse("actividad-calificaciones-plantilla", args=[(actividad or self.act1).id])
        )

    def subir(self, contenido, nombre="notas.xlsx", actividad=None, user=None):
        self.auth(user or self.doc1)
        archivo = SimpleUploadedFile(nombre, contenido, content_type=excel.CONTENT_TYPE)
        return self.client.post(
            reverse("actividad-calificaciones-importar", args=[(actividad or self.act1).id]),
            {"archivo": archivo},
            format="multipart",
        )

    def llenar(self, notas, contenido=None):
        """Toma la plantilla y escribe {estudiante_id o fila: valor} en la columna Nota."""
        wb = load_workbook(BytesIO(contenido or self.descargar().content))
        ws = wb[excel.HOJA]
        filas = {ws.cell(row=r, column=1).value: r for r in range(excel.FILA_ENCABEZADO + 1, ws.max_row + 1)}
        for clave, valor in notas.items():
            ws.cell(row=filas.get(clave, clave), column=3, value=valor)
        salida = BytesIO()
        wb.save(salida)
        return salida.getvalue()


class PlantillaTests(ExcelBase):
    def test_descarga_con_nomina_y_notas_actuales(self):
        self.guardar(self.doc1, self.act1, self.fracciones, {self.est1: "8.5"})
        res = self.descargar()
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res["Content-Type"], excel.CONTENT_TYPE)
        ws = load_workbook(BytesIO(res.content))[excel.HOJA]
        filas = [
            tuple(c.value for c in fila)
            for fila in ws.iter_rows(min_row=excel.FILA_ENCABEZADO + 1, max_col=3)
        ]
        # Solo el paralelo A, ordenado por apellido: Pérez, Vera.
        self.assertEqual(filas, [(self.est1.id, "Luis Pérez", 8.5), (self.est3.id, "Marta Vera", None)])
        self.assertTrue(ws.protection.sheet)

    def test_solo_el_docente_duenio(self):
        self.assertEqual(self.descargar(user=self.doc2).status_code, 404)
        self.assertEqual(self.descargar(user=self.rep1).status_code, 403)


class ImportarTests(ExcelBase):
    def test_vista_previa_sin_guardar(self):
        res = self.subir(self.llenar({self.est1.id: 9, self.est3.id: "6,5"}))
        self.assertEqual(res.status_code, 200, res.data)
        self.assertEqual(
            [(f["estudiante"], f["valor"]) for f in res.data["filas"]],
            [(self.est1.id, "9.00"), (self.est3.id, "6.50")],
        )
        self.assertEqual(res.data["errores"], [])
        self.assertFalse(Calificacion.objects.exists())  # no guarda nada

    def test_celdas_vacias_se_cuentan_sin_nota(self):
        res = self.subir(self.llenar({self.est1.id: 7}))
        self.assertEqual((len(res.data["filas"]), res.data["sin_nota"]), (1, 1))

    def test_errores_por_fila(self):
        contenido = self.llenar({self.est1.id: 12, self.est3.id: "abc"})
        wb = load_workbook(BytesIO(contenido))
        ws = wb[excel.HOJA]
        ws.append([self.est2.id, "Otro paralelo", 8])   # estudiante de otro paralelo
        ws.append([self.est1.id, "Duplicado", 8])        # repetido
        salida = BytesIO()
        wb.save(salida)
        res = self.subir(salida.getvalue())
        mensajes = " | ".join(e["mensaje"] for e in res.data["errores"])
        self.assertEqual(len(res.data["errores"]), 4, mensajes)
        self.assertIn("entre 0 y 10", mensajes)
        self.assertIn("no es un número", mensajes)
        self.assertIn("no pertenece a este paralelo", mensajes)
        self.assertIn("más de una vez", mensajes)

    def test_plantilla_de_otra_actividad_se_rechaza(self):
        otra = self.descargar(actividad=self.act2).content
        res = self.subir(otra, actividad=self.act1)
        self.assertEqual(res.status_code, 400)
        self.assertIn("otra actividad", str(res.data))

    def test_archivo_no_excel_se_rechaza(self):
        self.assertEqual(self.subir(b"hola", nombre="notas.csv").status_code, 400)
        self.assertEqual(self.subir(b"hola", nombre="notas.pdf").status_code, 400)
        self.assertEqual(self.subir(b"no es excel", nombre="notas.xlsx").status_code, 400)

    def test_columnas_cambiadas_usa_el_mapeo_inteligente(self):
        wb = load_workbook(BytesIO(self.llenar({self.est1.id: 9})))
        wb[excel.HOJA].cell(row=excel.FILA_ENCABEZADO, column=1, value="Código")
        salida = BytesIO()
        wb.save(salida)
        res = self.subir(salida.getvalue())
        self.assertEqual(res.status_code, 200, res.data)
        self.assertEqual(res.data["modo"], "mapeo")
        self.assertEqual([f["estudiante"] for f in res.data["filas"]], [self.est1.id])

    def test_otro_docente_no_puede_importar(self):
        res = self.subir(self.llenar({self.est1.id: 9}), user=self.doc2)
        self.assertEqual(res.status_code, 404)
