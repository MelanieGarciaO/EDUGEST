"""
Carga de notas con plantilla Excel fija (mejora 9).

1. El docente descarga la plantilla de la actividad: trae la nómina del paralelo
   (ID y nombre bloqueados) y una columna «Nota» editable, con las notas actuales.
2. Llena las notas en Excel y la sube.
3. El sistema valida el archivo y devuelve una VISTA PREVIA; no guarda nada. El docente
   revisa las notas en pantalla, elige el tema y pulsa «Guardar» (flujo normal).
"""

from decimal import Decimal, InvalidOperation
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation
from rest_framework.exceptions import ValidationError

TAMANO_MAXIMO = 2 * 1024 * 1024  # 2 MB
HOJA = "Notas"
HOJA_META = "_edugest"
ENCABEZADO = ("ID", "Estudiante", "Nota (0 a 10)")
FILA_ENCABEZADO = 5
CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def nombre_archivo(actividad):
    base = f"notas_{actividad.asignacion.asignatura.nombre}_{actividad.asignacion.paralelo.nombre_completo}"
    limpio = "".join(c if c.isalnum() else "_" for c in base)
    return f"{limpio}_act{actividad.id}.xlsx"


def generar_plantilla(actividad, estudiantes, notas):
    """
    estudiantes: queryset ordenado; notas: {estudiante_id: Decimal}. Devuelve bytes .xlsx.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = HOJA
    asignatura = actividad.asignacion.asignatura.nombre
    paralelo = actividad.asignacion.paralelo.nombre_completo

    ws["A1"] = "EduGest · Registro de calificaciones"
    ws["A1"].font = Font(bold=True, size=14, color="17406F")
    ws["A2"] = f"{asignatura} · {paralelo} · Entrega: {actividad.fecha_entrega:%d/%m/%Y}"
    ws["A3"] = f"Actividad: {actividad.descripcion[:200]}"
    ws["A3"].alignment = Alignment(wrap_text=False)
    ws["A4"] = "Escribe solo en la columna «Nota». Deja vacío a quien no vas a calificar."
    ws["A4"].font = Font(italic=True, color="5B6B80")

    cabecera = PatternFill("solid", fgColor="17406F")
    for col, texto in enumerate(ENCABEZADO, start=1):
        celda = ws.cell(row=FILA_ENCABEZADO, column=col, value=texto)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = cabecera

    editable = PatternFill("solid", fgColor="FFF8E1")
    fila = FILA_ENCABEZADO
    for fila, est in enumerate(estudiantes, start=FILA_ENCABEZADO + 1):
        ws.cell(row=fila, column=1, value=est.id)
        ws.cell(row=fila, column=2, value=est.full_name)
        nota = ws.cell(row=fila, column=3, value=float(notas[est.id]) if est.id in notas else None)
        nota.number_format = "0.00"
        nota.fill = editable
        nota.protection = Protection(locked=False)

    if fila > FILA_ENCABEZADO:
        dv = DataValidation(
            type="decimal",
            operator="between",
            formula1="0",
            formula2="10",
            allow_blank=True,
            showErrorMessage=True,
            errorTitle="Nota no válida",
            error="La nota debe estar entre 0 y 10.",
        )
        ws.add_data_validation(dv)
        dv.add(f"C{FILA_ENCABEZADO + 1}:C{fila}")

    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 42
    ws.column_dimensions["C"].width = 16
    ws.freeze_panes = ws.cell(row=FILA_ENCABEZADO + 1, column=1)
    # Hoja protegida (sin contraseña): solo la columna «Nota» se puede editar.
    ws.protection.sheet = True

    meta = wb.create_sheet(HOJA_META)
    meta["A1"] = "actividad"
    meta["B1"] = actividad.id
    meta.sheet_state = "hidden"

    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()


def _leer_nota(valor):
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return None
    try:
        nota = Decimal(str(valor).strip().replace(",", "."))
    except InvalidOperation:
        raise ValueError("no es un número") from None
    if nota.is_nan() or not Decimal("0") <= nota <= Decimal("10"):
        raise ValueError("debe estar entre 0 y 10")
    return nota.quantize(Decimal("0.01"))


def leer_plantilla(archivo, actividad, estudiantes):
    """
    Valida el archivo y devuelve la vista previa:
    {"filas": [{estudiante, estudiante_nombre, valor}], "errores": [{fila, mensaje}],
     "sin_nota": n}
    """
    if not archivo:
        raise ValidationError({"archivo": "Adjunta el archivo Excel de la plantilla."})
    if not archivo.name.lower().endswith(".xlsx"):
        raise ValidationError({"archivo": "El archivo debe ser .xlsx (la plantilla de EduGest)."})
    if archivo.size > TAMANO_MAXIMO:
        raise ValidationError({"archivo": "El archivo supera los 2 MB."})

    try:
        wb = load_workbook(archivo, data_only=True, read_only=False)
    except Exception:  # archivo dañado o que no es Excel
        raise ValidationError({"archivo": "No se pudo leer el archivo. Usa la plantilla descargada."}) from None

    if HOJA not in wb.sheetnames:
        raise ValidationError({"archivo": f"No se encontró la hoja «{HOJA}». Usa la plantilla descargada."})
    if HOJA_META in wb.sheetnames and wb[HOJA_META]["B1"].value not in (None, actividad.id):
        raise ValidationError(
            {"archivo": "Esta plantilla es de otra actividad. Descarga la de esta actividad."}
        )

    ws = wb[HOJA]
    encabezado = tuple(
        str(ws.cell(row=FILA_ENCABEZADO, column=c).value or "").strip() for c in (1, 2, 3)
    )
    if encabezado[0] != "ID" or not encabezado[2].startswith("Nota"):
        raise ValidationError(
            {"archivo": "El formato no coincide con la plantilla. No cambies las columnas."}
        )

    nombres = {e.id: e.full_name for e in estudiantes}
    filas, errores, vistos, sin_nota = [], [], set(), 0
    for n, (id_celda, _nombre, nota_celda) in enumerate(
        ws.iter_rows(min_row=FILA_ENCABEZADO + 1, max_col=3, values_only=True),
        start=FILA_ENCABEZADO + 1,
    ):
        if id_celda in (None, "") and nota_celda in (None, ""):
            continue
        try:
            est_id = int(id_celda)
        except (TypeError, ValueError):
            errores.append({"fila": n, "mensaje": "ID de estudiante no válido."})
            continue
        if est_id not in nombres:
            errores.append({"fila": n, "mensaje": "El estudiante no pertenece a este paralelo."})
            continue
        if est_id in vistos:
            errores.append({"fila": n, "mensaje": f"{nombres[est_id]} aparece más de una vez."})
            continue
        vistos.add(est_id)
        try:
            valor = _leer_nota(nota_celda)
        except ValueError as exc:
            errores.append({"fila": n, "mensaje": f"Nota de {nombres[est_id]}: {exc}."})
            continue
        if valor is None:
            sin_nota += 1
            continue
        filas.append({"estudiante": est_id, "estudiante_nombre": nombres[est_id], "valor": str(valor)})

    return {"filas": filas, "errores": errores, "sin_nota": sin_nota}
